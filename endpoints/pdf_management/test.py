from docx import Document
from fpdf import FPDF

# Open the word document
document = Document("original.docx")

# Create a new pdf document
pdf = FPDF()
pdf.add_font('Arial Unicode', '', 'Arial Unicode MS Font.ttf', uni=True)
pdf.set_font("Arial Unicode", size=12)
pdf.add_page()

# Iterate through the paragraphs in the word document
for para in document.paragraphs:
    # Extract the text and formatting of the paragraph
    text = para.text
    text = text.replace("offer_to", "new text")
    alignment = para.alignment
    # Split the text into lines
    lines = text.splitlines()
    pdf.set_xy(0, pdf.get_y()+10)
    for line in lines:
        if alignment == 'center':
            pdf.cell(0, 10, txt=line, align='C')
        elif alignment == 'right':
            pdf.cell(0, 10, txt=line, align='R')
        else:
            pdf.cell(0, 10, txt=line, align='L')

pdf.output("modified.pdf")
